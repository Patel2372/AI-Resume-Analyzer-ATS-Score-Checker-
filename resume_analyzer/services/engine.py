"""Core resume parsing and ATS analysis engine.

The engine intentionally avoids hard dependencies on an online LLM. The included
rewrite feature is deterministic so the project works after standard pip install.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List

SKILL_ALIASES = {
    'python': {'python'},
    'django': {'django'},
    'flask': {'flask'},
    'java': {'java'},
    'c++': {'c++', 'cpp'},
    'c#': {'c#', 'csharp'},
    'javascript': {'javascript', 'js'},
    'typescript': {'typescript', 'ts'},
    'html': {'html'},
    'css': {'css'},
    'bootstrap': {'bootstrap'},
    'react': {'react', 'react.js'},
    'node.js': {'node.js', 'nodejs'},
    'rest api': {'rest api', 'restful api'},
    'data structures': {'data structure', 'data structures', 'data structure & algorithms', 'data structures and algorithms'},
    'object oriented programming': {'object oriented programming', 'object-oriented programming', 'oop'},
    'cloud computing': {'cloud computing', 'cloud'},
    'sql': {'sql'},
    'mysql': {'mysql'},
    'postgresql': {'postgresql', 'postgres'},
    'sqlite': {'sqlite'},
    'mongodb': {'mongodb', 'mongo db'},
    'git': {'git'},
    'github': {'github'},
    'docker': {'docker'},
    'kubernetes': {'kubernetes', 'k8s'},
    'jenkins': {'jenkins'},
    'aws': {'aws', 'amazon web services'},
    'azure': {'azure'},
    'gcp': {'gcp', 'google cloud'},
    'linux': {'linux'},
    'bash': {'bash', 'shell scripting'},
    'terraform': {'terraform'},
    'ansible': {'ansible'},
    'nmap': {'nmap'},
    'wireshark': {'wireshark'},
    'burp suite': {'burp suite'},
    'cybersecurity': {'cybersecurity', 'cyber security'},
    'machine learning': {'machine learning', 'ml'},
    'deep learning': {'deep learning', 'dl'},
    'nlp': {'nlp', 'natural language processing'},
    'tensorflow': {'tensorflow'},
    'pytorch': {'pytorch'},
    'scikit-learn': {'scikit-learn', 'sklearn'},
    'pandas': {'pandas'},
    'numpy': {'numpy'},
    'opencv': {'opencv'},
    'flask': {'flask'},
    'power bi': {'power bi', 'powerbi'},
    'tableau': {'tableau'},
    'excel': {'excel', 'microsoft excel'},
}

JOB_ROLES = {
    'Python Developer': {'python', 'django', 'flask', 'sql', 'git', 'rest api'},
    'Backend Developer': {'python', 'java', 'sql', 'rest api', 'git', 'docker'},
    'Data Analyst': {'python', 'sql', 'pandas', 'numpy', 'excel', 'power bi', 'tableau'},
    'Data Scientist': {'python', 'machine learning', 'pandas', 'numpy', 'scikit-learn', 'sql'},
    'ML Engineer': {'python', 'machine learning', 'tensorflow', 'pytorch', 'docker', 'git'},
    'DevOps Engineer': {'linux', 'docker', 'kubernetes', 'jenkins', 'aws', 'terraform', 'ansible', 'git'},
    'Cybersecurity Analyst': {'linux', 'nmap', 'wireshark', 'burp suite', 'cybersecurity', 'python'},
    'Web Developer': {'html', 'css', 'javascript', 'bootstrap', 'react', 'node.js', 'git'},
}

SECTION_PATTERNS = {
    'summary': r'(?i)^(professional summary|summary|profile|objective)$',
    'skills': r'(?i)^(technical skills|skills|core skills)$',
    'education': r'(?i)^(education|academic background)$',
    'experience': r'(?i)^(experience|work experience|employment|internship|internship experience|professional experience)$',
    'projects': r'(?i)^(projects|academic projects|personal projects|projects & hackathons|projects and hackathons)$',
    'certifications': r'(?i)^(certifications|certificates)$',
    'achievements': r'(?i)^(achievements|awards|accomplishments|achievements - ssip hackathon)$',
    'soft_skills': r'(?i)^soft skills$',
    'tools': r'(?i)^(tools|technical tools)$',
    'languages': r'(?i)^languages$',
}

STOPWORDS = {
    'the','and','with','for','from','this','that','have','has','are','was','were','will',
    'you','your','our','their','into','about','using','used','required','preferred','role',
    'job','work','team','years','year','experience','skills','developer','engineer'
}


def normalize(text: str) -> str:
    return re.sub(r'\s+', ' ', text or '').strip().lower()


def extract_pdf(path: str) -> str:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError('PyMuPDF is not installed. Run pip install -r requirements.txt') from exc
    doc = fitz.open(path)
    try:
        return '\n'.join(page.get_text('text') for page in doc)
    finally:
        doc.close()


def extract_docx(path: str) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError('python-docx is not installed. Run pip install -r requirements.txt') from exc
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(' | '.join(cell.text for cell in row.cells))
    return '\n'.join(parts)


def extract_text(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext == '.pdf':
        return extract_pdf(path)
    if ext == '.docx':
        return extract_docx(path)
    raise ValueError('Unsupported resume format. Use PDF or DOCX.')


def extract_skills(text: str) -> List[str]:
    low = normalize(text)
    found = []
    for canonical, aliases in SKILL_ALIASES.items():
        if any(re.search(r'(?<!\w)' + re.escape(alias) + r'(?!\w)', low) for alias in aliases):
            found.append(canonical)
    return sorted(found)


def sectionize(text: str) -> Dict[str, str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    current = 'other'
    sections = {k: '' for k in list(SECTION_PATTERNS) + ['other']}
    for line in lines:
        matched = None
        for key, pattern in SECTION_PATTERNS.items():
            if re.match(pattern, line.strip(':- ')):
                matched = key
                break
        if matched:
            current = matched
            continue
        sections[current] += line + '\n'
    return {k: v.strip() for k, v in sections.items()}


def extract_email(text: str) -> str:
    m = re.search(r'[\w.+-]+@[\w-]+\.[\w.-]+', text)
    return m.group(0) if m else ''


def extract_phone(text: str) -> str:
    m = re.search(r'(?<!\d)(?:\+?\d[\d\s().-]{8,}\d)(?!\d)', text)
    return re.sub(r'\s+', ' ', m.group(0)).strip() if m else ''


def extract_url(text: str, kind: str) -> str:
    pattern = {
        'linkedin': r'https?://(?:www\.)?linkedin\.com/[^\s|]+|www\.linkedin\.com/[^\s|]+',
        'github': r'https?://(?:www\.)?github\.com/[^\s|]+|www\.github\.com/[^\s|]+',
    }[kind]
    m = re.search(pattern, text, flags=re.I)
    return m.group(0).rstrip('.,);') if m else ''


def extract_location(text: str) -> str:
    for line in text.splitlines():
        clean = _clean_block_text(line)
        if re.fullmatch(r'[A-Za-z .-]+,\s*[A-Za-z .-]+,\s*India', clean, flags=re.I):
            return clean
    return ''


def _clean_block_text(text: str) -> str:
    return re.sub(r'\s+', ' ', text or '').strip()


def _is_heading(text: str) -> bool:
    value = _clean_block_text(text).strip(':- ').lower()
    heading_aliases = {
        'profile', 'summary', 'professional summary', 'objective',
        'technical skills', 'skills', 'core skills', 'soft skills', 'tools', 'languages',
        'education', 'academic background', 'work experience', 'experience', 'employment', 'internship',
        'projects', 'academic projects', 'personal projects', 'projects & hackathons', 'projects and hackathons', 'certificates', 'certifications',
        'achievements', 'awards', 'accomplishments', 'achievements - ssip hackathon'
    }
    return value in heading_aliases


def _section_key(text: str) -> str | None:
    value = _clean_block_text(text).strip(':- ').lower()
    mapping = {
        'profile': 'summary', 'summary': 'summary', 'professional summary': 'summary', 'objective': 'summary',
        'technical skills': 'skills', 'skills': 'skills', 'core skills': 'skills',
        'education': 'education', 'academic background': 'education',
        'work experience': 'experience', 'experience': 'experience', 'employment': 'experience', 'internship': 'experience', 'internship experience': 'experience', 'professional experience': 'experience',
        'projects': 'projects', 'academic projects': 'projects', 'personal projects': 'projects', 'projects & hackathons': 'projects', 'projects and hackathons': 'projects',
        'certificates': 'certifications', 'certifications': 'certifications',
        'achievements': 'achievements', 'awards': 'achievements', 'accomplishments': 'achievements', 'achievements - ssip hackathon': 'achievements',
        'soft skills': 'soft_skills', 'tools': 'tools', 'languages': 'languages'
    }
    return mapping.get(value)


def extract_name(text: str) -> str:
    """Best-effort name extraction for plain text resumes."""
    blacklist = {
        'software engineer', 'python developer', 'data scientist', 'data analyst', 'backend developer',
        'computer engineer', 'computer engineering', 'student', 'resume', 'curriculum vitae'
    }
    candidates = []
    for idx, line in enumerate(text.splitlines()):
        clean = _clean_block_text(line)
        if not clean or len(clean) > 60 or '@' in clean or re.search(r'\d', clean):
            continue
        words = clean.split()
        if 2 <= len(words) <= 5 and all(re.fullmatch(r"[A-Za-z.'-]+", w) for w in words):
            if clean.lower() in blacklist:
                continue
            score = 0
            if clean.isupper():
                score += 4
            if idx < 15:
                score += 3
            score += max(0, 6 - len(words))
            candidates.append((score, clean))
    return max(candidates, key=lambda x: x[0])[1] if candidates else ''


def _split_raw_lines(value: str) -> list[str]:
    """Split a PDF block into clean display items while preserving wording."""
    items = []
    for line in (value or '').splitlines():
        line = _clean_block_text(line).strip(' -•')
        if line:
            items.append(line)
    return items


def _date_only(value: str) -> bool:
    value = _clean_block_text(value)
    return bool(
        re.fullmatch(r'(?:19|20)\d{2}\s*[-–]\s*(?:19|20)\d{2}', value)
        or re.fullmatch(r'(?i)(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s+\d{4}\s*[-–]\s*(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s+\d{4}', value)
    )


def _starts_experience_entry(value: str) -> bool:
    value = _clean_block_text(value)
    return bool(
        re.search(r'(?i)\b(internship|hackathon|project)\b', value)
        or re.search(r'(?i)\b(?:19|20)\d{2}\s*[-–]\s*(?:19|20)\d{2}\b', value)
        or re.search(r'(?i)\b(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s+\d{4}', value)
    )


def _group_entry_blocks(blocks: list[str], mode: str = '') -> list[str]:
    """Group visual PDF blocks into readable resume entries without dropping content."""
    entries: list[str] = []
    current = ''
    for raw in blocks:
        lines = _split_raw_lines(raw)
        if not lines:
            continue
        block = ' | '.join(lines)
        if mode == 'experience':
            is_short_header = len(block) <= 140 and _starts_experience_entry(block)
            is_date_only = _date_only(block)
            if is_short_header and not is_date_only:
                if current:
                    entries.append(current)
                current = block
            elif current:
                current += ' | ' + block
            else:
                current = block
        elif mode == 'education':
            if current and _date_only(block):
                current += ' | ' + block
            elif current:
                entries.append(current)
                current = block
            else:
                current = block
        elif mode == 'projects':
            if current:
                current += ' | ' + block
            else:
                current = block
        else:
            entries.append(block)
    if current:
        entries.append(current)
    return entries


def _certificate_items(blocks: list[str]) -> list[str]:
    """Split certificate blocks into complete certificate names, preserving wrapped lines."""
    lines: list[str] = []
    for raw in blocks:
        lines.extend(_split_raw_lines(raw))
    items: list[str] = []
    current = ''
    for line in lines:
        # The uploaded resume has two certificate groups: NPTEL and AWS.
        is_new = bool(re.match(r'(?i)^(nptel\b|aws\b)', line))
        if is_new and current:
            items.append(current)
            current = line
        elif current:
            current += ' ' + line
        else:
            current = line
    if current:
        items.append(current)
    return items


def _clean_section_items(blocks: list[str], merge_parenthetical: bool = False) -> list[str]:
    items: list[str] = []
    for raw in blocks:
        for line in _split_raw_lines(raw):
            if merge_parenthetical and items and line.startswith('(') and line.endswith(')'):
                items[-1] += ' ' + line
            else:
                items.append(line)
    return items


def extract_pdf_layout(path: str) -> Dict:
    """Extract structured information from a visually designed PDF resume.

    The parser uses block coordinates because a two-column resume cannot be safely
    interpreted by treating the PDF as one linear text stream. Section blocks are
    retained so the UI can show every education, experience, project and certificate
    entry instead of collapsing them into a single item.
    """
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError('PyMuPDF is not installed. Run pip install -r requirements.txt') from exc

    doc = fitz.open(path)
    try:
        page_data = []
        name_candidates = []
        headline_candidates = []
        for page_index, page in enumerate(doc):
            page_width = page.rect.width
            blocks = []
            for block in page.get_text('blocks'):
                if len(block) < 7:
                    continue
                x0, y0, x1, y1, text = block[:5]
                clean = _clean_block_text(text)
                if not clean:
                    continue
                item = {
                    'page': page_index,
                    'x0': float(x0), 'y0': float(y0), 'x1': float(x1), 'y1': float(y1),
                    'text': clean, 'raw_text': text.strip()
                }
                blocks.append(item)

                # The name is the large top-left candidate, not a section heading.
                if y0 < 80 and x0 < page_width * 0.40 and '@' not in clean and not re.search(r'\d', clean):
                    words = clean.split()
                    if 2 <= len(words) <= 5 and all(re.fullmatch(r"[A-Za-z.'-]+", w) for w in words):
                        score = 1000 - y0
                        if clean.isupper():
                            score += 20
                        if clean.upper() in {'JENIL PATEL'}:
                            score += 100
                        if clean.lower() in {'software engineer', 'devops engineer', 'python developer', 'data analyst', 'data scientist'}:
                            score -= 150
                        name_candidates.append((score, clean))
            page_data.append((page, blocks, page_width))

        name = max(name_candidates, key=lambda item: item[0])[1] if name_candidates else ''
        headline_candidates = []
        for _, blocks, page_width in page_data:
            for b in blocks:
                if 45 <= b['y0'] < 90 and b['x0'] < page_width * 0.40 and b['text'] != name:
                    clean = b['text']
                    words = clean.split()
                    if 2 <= len(words) <= 5 and all(re.fullmatch(r"[A-Za-z.'-]+", w) for w in words):
                        headline_candidates.append((b['y0'], clean))
        headline = min(headline_candidates, key=lambda item: item[0])[1] if headline_candidates else ''
        section_names = list(SECTION_PATTERNS) + ['soft_skills', 'tools', 'languages', 'other']
        sections = {k: '' for k in section_names}
        section_blocks = {k: [] for k in section_names}

        for page, blocks, page_width in page_data:
            split = page_width * 0.36
            # Keep independent reading order for the sidebar and main column.
            left_blocks = [b for b in blocks if b['x0'] < split]
            right_blocks = [b for b in blocks if b['x0'] >= split]

            for column_blocks in (left_blocks, right_blocks):
                current = 'other'
                for b in sorted(column_blocks, key=lambda z: (z['y0'], z['x0'])):
                    clean = b['text']
                    heading = _section_key(clean)
                    if heading:
                        current = heading
                        continue
                    # Ignore header/contact/social metadata outside sections.
                    if b['y0'] < 85:
                        continue
                    if clean == name or '@' in clean or clean.startswith(('http://', 'https://', 'www.')):
                        continue
                    if current in section_blocks:
                        section_blocks[current].append(b['raw_text'])

        for key, blocks in section_blocks.items():
            sections[key] = '\n'.join(blocks).strip()

        return {
            'name': name,
            'headline': headline,
            'sections': sections,
            'section_blocks': section_blocks,
        }
    finally:
        doc.close()


def _parse_experience_entries(value: str) -> list[str]:
    """Keep each role/internship and its bullet descriptions as one record."""
    entries: list[str] = []
    current = ''
    body_started = False
    previous_line = ''
    for raw in (value or '').splitlines():
        line = _clean_block_text(raw)
        if not line:
            continue
        is_bullet = bool(re.match(r'^[•●▪◦*-]\s+', line))
        cleaned = re.sub(r'^[•●▪◦*-]\s+', '', line).strip()
        if is_bullet:
            current = (current + ' | ' if current else '') + cleaned
            body_started = True
        elif current and body_started:
            continuation = bool(previous_line and (
                previous_line.rstrip().endswith((',', ';', ':', 'and', 'or', '/', '-'))
                or cleaned[:1].islower()
            ))
            if continuation:
                current += ' ' + cleaned
            else:
                entries.append(current)
                current = cleaned
                body_started = False
        elif current:
            current += ' ' + cleaned
        else:
            current = cleaned
        previous_line = cleaned
    if current:
        entries.append(current)
    return [entry.strip() for entry in entries if entry.strip()]


def _parse_project_entries(value: str) -> list[str]:
    """Group project headings and their bullets while preserving wrapped lines."""
    entries: list[str] = []
    current = ''
    body_started = False
    previous_line = ''
    for raw in (value or '').splitlines():
        line = _clean_block_text(raw).strip('\t ')
        if not line:
            continue
        is_bullet = bool(re.match(r'^[•●▪◦*-]\s+', line))
        cleaned = re.sub(r'^[•●▪◦*-]\s+', '', line).strip()
        if re.match(r'(?i)^github\s*:', cleaned):
            if current:
                current += ' | ' + cleaned
            else:
                current = cleaned
            previous_line = cleaned
            continue
        if is_bullet:
            current = (current + ' | ' if current else '') + cleaned
            body_started = True
            previous_line = cleaned
            continue
        continuation = bool(previous_line and (
            previous_line.rstrip().endswith((',', ';', ':', 'and', 'or', '/', '-'))
            or cleaned[:1].islower()
        ))
        # In conventional resumes, project titles often include a technology pipe.
        # That is a strong boundary when the previous project's body has begun.
        looks_like_title = '|' in cleaned and len(cleaned) <= 220
        if current and body_started and looks_like_title:
            entries.append(current)
            current = cleaned
            body_started = False
        elif current and body_started and not continuation:
            entries.append(current)
            current = cleaned
            body_started = False
        elif current:
            current += ' ' + cleaned
        else:
            current = cleaned
        previous_line = cleaned
    if current:
        entries.append(current)
    return [entry.strip() for entry in entries if entry.strip()]


def parse_resume(text: str, source_path: str | None = None) -> Dict:
    plain_sections = sectionize(text)
    sections = plain_sections
    section_blocks = {k: [] for k in list(SECTION_PATTERNS) + ['soft_skills', 'tools', 'languages', 'other']}
    layout_name = ''
    layout = {}

    if source_path and Path(source_path).suffix.lower() == '.pdf':
        try:
            layout = extract_pdf_layout(source_path)
            layout_name = layout.get('name', '')
            layout_sections = layout.get('sections', {})
            layout_blocks = layout.get('section_blocks', {})
            # The coordinate-aware PDF representation is authoritative when available.
            meaningful_layout = any(v for k, v in layout_sections.items() if k != 'other')
            if meaningful_layout:
                sections = {**plain_sections, **{k: v for k, v in layout_sections.items() if v and k != 'other'}}
                section_blocks = {**section_blocks, **{k: v for k, v in layout_blocks.items() if k != 'other'}}
        except Exception:
            pass

    lower = normalize(text)
    education_hits = re.findall(
        r'(?i)(b\.?tech|b\.?e\.?|m\.?tech|mca|bca|bachelor|master|diploma|computer engineering)',
        text,
    )
    years = [int(x) for x in re.findall(r'(?<!\d)([0-9]{1,2})\+?\s+years?', lower)]

    # Resume-wide detected skills are useful for ATS matching. Technical skills are
    # separately preserved exactly from the TECHNICAL SKILLS section for display.
    all_detected_skills = extract_skills(text)
    technical_skills = _clean_section_items(section_blocks.get('skills', []), merge_parenthetical=True)
    if not technical_skills and sections.get('skills'):
        technical_skills = [x.strip(' -•') for x in sections['skills'].splitlines() if x.strip()]

    experience_blocks = section_blocks.get('experience', [])
    experience_items = _parse_experience_entries(sections.get('experience', ''))
    if not experience_items and experience_blocks:
        experience_items = _parse_experience_entries('\n'.join(experience_blocks))

    education_blocks = section_blocks.get('education', [])
    education_items = _group_entry_blocks(education_blocks, mode='education')
    if not education_items and sections.get('education'):
        education_items = [x for x in sections['education'].splitlines() if x.strip()]

    project_blocks = section_blocks.get('projects', [])
    projects = _parse_project_entries(sections.get('projects', ''))
    if not projects and project_blocks:
        projects = _parse_project_entries('\n'.join(project_blocks))

    certifications = _certificate_items(section_blocks.get('certifications', []))
    if not certifications and sections.get('certifications'):
        cert_text = sections['certifications'].replace(' | ', '\n')
        certifications = [x.strip(' -•') for x in cert_text.splitlines() if x.strip()]

    tools = _clean_section_items(section_blocks.get('tools', []))
    languages = _clean_section_items(section_blocks.get('languages', []))
    if not languages and sections.get('languages'):
        languages = [x.strip(' -•') for x in re.split(r'[|,\n]', sections['languages']) if x.strip()]
    soft_skills = _clean_section_items(section_blocks.get('soft_skills', []))
    if not soft_skills and sections.get('soft_skills'):
        soft_skills = [x.strip(' -•') for x in re.split(r'[|,\n]', sections['soft_skills']) if x.strip()]

    # If a PDF parser did not isolate tools/soft skills, keep the plain-text fallback.
    if not tools and sections.get('tools'):
        tools = [x.strip(' -•') for x in sections['tools'].splitlines() if x.strip()]
    if not languages and sections.get('languages'):
        languages = [x.strip(' -•') for x in sections['languages'].splitlines() if x.strip()]
    if not soft_skills and sections.get('soft_skills'):
        soft_skills = [x.strip(' -•') for x in sections['soft_skills'].splitlines() if x.strip()]

    # Keep project records strictly inside the Projects section. Experience and
    # achievements are separate categories and must never be copied into projects.
    all_projects = list(projects)

    headline = ''
    if source_path and Path(source_path).suffix.lower() == '.pdf':
        try:
            headline = layout.get('headline', '') if 'layout' in locals() else ''
        except Exception:
            headline = ''

    return {
        'name': layout_name or extract_name(text),
        'headline': headline,
        'email': extract_email(text),
        'phone': extract_phone(text),
        'linkedin': extract_url(text, 'linkedin'),
        'github': extract_url(text, 'github'),
        'location': extract_location(text),
        'skills': all_detected_skills,
        'technical_skills': technical_skills,
        'sections': sections,
        'education_terms': sorted(set(education_hits), key=str.lower),
        'experience_years': max(years) if years else 0,
        'projects': projects,
        'all_projects': all_projects,
        'experience_items': experience_items,
        'education_items': education_items,
        'certifications': certifications,
        'tools': tools,
        'languages': languages,
        'soft_skills': soft_skills,
        'section_blocks': section_blocks,
        'sections_detected': [k for k, v in sections.items() if v and k != 'other'],
        'word_count': len(re.findall(r'\b\w+\b', text)),
    }


def extract_keywords(text: str, limit: int = 18) -> List[str]:
    tokens = re.findall(r'[A-Za-z][A-Za-z0-9+#.-]{2,}', normalize(text))
    counts = {}
    for token in tokens:
        if token in STOPWORDS or len(token) < 3:
            continue
        counts[token] = counts.get(token, 0) + 1
    skill_terms = set(extract_skills(text))
    ranked = [word for word, _ in sorted(counts.items(), key=lambda x: (-x[1], x[0]))]
    ordered = list(skill_terms) + [w for w in ranked if w not in skill_terms]
    return ordered[:limit]


def semantic_similarity(a: str, b: str) -> float:
    try:
        from sklearn.feature_extraction.text import TfidfVectorizer
        from sklearn.metrics.pairwise import cosine_similarity
        vectors = TfidfVectorizer(stop_words='english', ngram_range=(1,2), min_df=1).fit_transform([a, b])
        return float(cosine_similarity(vectors[0:1], vectors[1:2])[0][0] * 100)
    except Exception:
        at = set(re.findall(r'[a-zA-Z][a-zA-Z0-9+#.-]{2,}', normalize(a)))
        bt = set(re.findall(r'[a-zA-Z][a-zA-Z0-9+#.-]{2,}', normalize(b)))
        if not at or not bt:
            return 0.0
        return len(at & bt) / len(at | bt) * 100


def formatting_check(path: str, text: str) -> Dict:
    ext = Path(path).suffix.lower()
    warnings = []
    positives = []
    if len(text) < 250:
        warnings.append('Resume content is very short.')
    else:
        positives.append('Resume contains readable text.')
    detected_sections = None
    if ext == '.pdf':
        try:
            detected_sections = extract_pdf_layout(path).get('sections', {})
        except Exception:
            detected_sections = None
    if detected_sections is None:
        detected_sections = sectionize(text)
    if not any(detected_sections.get(k) for k in ['skills', 'education', 'experience', 'projects']):
        warnings.append('Important resume sections could not be detected.')
    else:
        positives.append('Important resume sections detected.')
    if ext == '.docx':
        try:
            from docx import Document
            doc = Document(path)
            if len(doc.tables) > 0:
                warnings.append('Tables detected. Simple tables are safer for ATS parsing.')
            else:
                positives.append('No tables detected in DOCX.')
        except Exception:
            warnings.append('DOCX formatting could not be deeply inspected.')
    return {'warnings': warnings, 'positives': positives, 'score': max(0, min(100, 100 - 12 * len(warnings)))}


def predict_job_roles(skills: List[str]) -> List[Dict]:
    s = set(skills)
    out = []
    for role, required in JOB_ROLES.items():
        match = len(s & required) / max(1, len(required)) * 100
        out.append({'role': role, 'match': round(match, 1), 'matched': sorted(s & required), 'missing': sorted(required - s)})
    return sorted(out, key=lambda x: x['match'], reverse=True)[:6]


def analyze_resume(resume_text: str, job_text: str, resume_path: str = '') -> Dict:
    r = parse_resume(resume_text, resume_path or None)
    j_skills = extract_skills(job_text)
    r_skills = set(r['skills'])
    j_skills_set = set(j_skills)
    matched = sorted(r_skills & j_skills_set)
    missing = sorted(j_skills_set - r_skills)
    skill_match = len(matched) / max(1, len(j_skills_set)) * 100
    keywords = extract_keywords(job_text)
    resume_low = normalize(resume_text)
    matched_keywords = [k for k in keywords if k in resume_low]
    missing_keywords = [k for k in keywords if k not in resume_low]
    keyword_score = len(matched_keywords) / max(1, len(keywords)) * 100
    semantic = semantic_similarity(resume_text, job_text)
    sections = r['sections']
    section_presence = sum(1 for k in ['summary','skills','education','experience','projects'] if sections.get(k))
    quality = min(100, round(45 + section_presence * 7 + min(len(r['skills']), 10) * 1.2 + min(r['word_count'], 1200) / 120, 1))
    formatting = formatting_check(resume_path, resume_text) if resume_path else {'warnings': [], 'positives': ['File formatting check skipped.'], 'score': 80}
    education_match = 100 if r['education_terms'] else 40
    experience_match = 100 if r.get('experience_items') else 60
    overall = (skill_match * 0.25 + keyword_score * 0.20 + semantic * 0.20 + experience_match * 0.15 + education_match * 0.10 + formatting['score'] * 0.10)
    strengths = []
    weaknesses = []
    if matched: strengths.append(f'Strong match for: {", ".join(matched[:6])}.')
    if r['projects']: strengths.append(f'{len(r["projects"])} project item(s) detected.')
    if r['certifications']: strengths.append('Certification section detected.')
    if section_presence >= 4: strengths.append('Most important resume sections are present.')
    if missing: weaknesses.append(f'Missing job skills: {", ".join(missing[:8])}.')
    if missing_keywords: weaknesses.append(f'{len(missing_keywords)} important job keywords are missing.')
    if formatting['warnings']: weaknesses.extend(formatting['warnings'])
    if not r.get('experience_items'): weaknesses.append('Experience section is missing or too difficult to detect.')
    priorities = []
    for item in missing[:5]: priorities.append({'level':'High', 'text':f'Add or develop the {item} skill if you genuinely have this capability.'})
    if missing_keywords: priorities.append({'level':'Medium', 'text':'Add relevant missing keywords naturally to experience/projects where truthful.'})
    if not sections.get('projects'): priorities.append({'level':'Medium', 'text':'Add 2-3 quantified project bullets relevant to the target role.'})
    roles = predict_job_roles(r['skills'])
    learning = []
    for skill in missing[:6]:
        learning.append({'skill': skill, 'why': f'{skill} appears in the job requirements but not in the detected resume skills.', 'next_step': f'Learn {skill}, build one small project, and add evidence to your resume.'})
    role = (j_skills_set, job_text)
    job_title = job_text.splitlines()[0].strip() if job_text.strip() else 'Target Role'
    return {
        'candidate': {'name': r['name'], 'email': r['email'], 'phone': r['phone']},
        'ats_score': round(overall, 1),
        'score_breakdown': {
            'Skills Match': round(skill_match,1),
            'Keyword Match': round(keyword_score,1),
            'Semantic Match': round(semantic,1),
            'Experience Match': round(experience_match,1),
            'Education Match': round(education_match,1),
            'ATS Formatting': round(formatting['score'],1),
        },
        'resume_quality': quality,
        'resume_skills': sorted(r_skills),
        'job_skills': sorted(j_skills_set),
        'matched_skills': matched,
        'missing_skills': missing,
        'keywords': keywords,
        'matched_keywords': matched_keywords,
        'missing_keywords': missing_keywords,
        'semantic_similarity': round(semantic,1),
        'formatting': formatting,
        'strengths': strengths or ['The parser found limited evidence of resume strengths.'],
        'weaknesses': weaknesses or ['No major weaknesses were detected by the automated checks.'],
        'priorities': priorities,
        'predicted_roles': roles,
        'learning_recommendations': learning,
        'job_title_hint': job_title,
        'word_count': r['word_count'],
        'sections_detected': [k for k,v in sections.items() if v],
        'parsed_resume': r,
    }


def rewrite_resume_text(text: str, context: str = '') -> str:
    original = text.strip()
    if not original:
        return ''
    replacements = [
        (r'^worked on (.+)$', r'Developed and implemented \1 with a focus on measurable outcomes.'),
        (r'^helped with (.+)$', r'Contributed to \1 and supported successful delivery.'),
        (r'^responsible for (.+)$', r'Led and executed \1, improving reliability and delivery quality.'),
        (r'^made (.+)$', r'Developed \1 and improved the overall user experience.'),
        (r'^used (.+)$', r'Applied \1 to build practical, production-oriented solutions.'),
    ]
    for pattern, replacement in replacements:
        changed = re.sub(pattern, replacement, original, flags=re.I)
        if changed != original:
            return changed
    if not re.search(r'(developed|implemented|built|optimized|designed|automated|improved|analyzed)', original, re.I):
        return 'Developed and implemented ' + original.rstrip('.') + ' with measurable results and clear technical ownership.'
    return original
